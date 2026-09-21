import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
import django
django.setup()

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas
from portfolio.models import Profile, Project

OUT = Path("output")
OUT.mkdir(exist_ok=True)
PPTX_PATH = OUT / "Rashid_Hassan_Ally_Portfolio_Editable.pptx"
PDF_PATH = OUT / "Rashid_Hassan_Ally_Portfolio.pdf"

BG, CARD, TEXT, MUTED, PINK, PURPLE = "0D0D0F", "16161A", "F5F5F7", "9A9AA2", "FF4D6D", "7C5CFF"


def rgb(value):
    return RGBColor.from_string(value)


def clean(value):
    return (value or "").replace("�", "—")


def fit_box(image_path, x, y, width, height):
    with Image.open(image_path) as image:
        ratio = image.width / image.height
    box_ratio = width / height
    if ratio > box_ratio:
        final_w, final_h = width, width / ratio
    else:
        final_w, final_h = height * ratio, height
    return x + (width - final_w) / 2, y + (height - final_h) / 2, final_w, final_h


profile = Profile.objects.first()
if not profile:
    raise SystemExit("Add a Portfolio Profile in Django admin before generating documents.")
projects = list(Project.objects.filter(featured=True))


def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(BG)


def add_text(slide, text, x, y, w, h, size=18, color=TEXT, bold=False, font="Aptos", align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = clean(text)
    p.font.name, p.font.size, p.font.bold, p.font.color.rgb = font, Pt(size), bold, rgb(color)
    if align:
        p.alignment = align
    return box


def card(slide, x, y, w, h):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(CARD)
    shape.line.color.rgb = rgb("29292F")
    return shape


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

# Cover
slide = prs.slides.add_slide(blank); set_background(slide)
shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.5), Inches(-1.5), Inches(6), Inches(6))
shape.fill.solid(); shape.fill.fore_color.rgb = rgb(PURPLE); shape.fill.transparency = 45; shape.line.fill.background()
add_text(slide, profile.name.upper(), .8, 1.7, 9.6, .8, 31, TEXT, True)
add_text(slide, clean(profile.title).upper(), .85, 2.65, 7, .35, 13, PINK, True)
add_text(slide, "PORTFOLIO 2026", .85, 3.3, 5, .3, 12, MUTED)
add_text(slide, "Bold visuals. Clear ideas. Meaningful design.", .85, 5.65, 7.8, .5, 19, TEXT, False)
if profile.about_image:
    img = Path(profile.about_image.path)
    if img.exists():
        x, y, w, h = fit_box(img, 10.0, 3.9, 2.2, 2.2)
        slide.shapes.add_picture(str(img), Inches(x), Inches(y), Inches(w), Inches(h))

# About
slide = prs.slides.add_slide(blank); set_background(slide)
add_text(slide, "ABOUT ME", .8, .65, 5, .5, 24, TEXT, True)
add_text(slide, clean(profile.about_text), .85, 1.55, 6.3, 1.5, 16, MUTED)
add_text(slide, clean(profile.about_detail), .85, 3.1, 6.3, 1.6, 14, MUTED)
add_text(slide, "SKILLS", .85, 5.25, 2, .3, 12, PINK, True)
for index, skill in enumerate([x.strip() for x in profile.skills.split(",") if x.strip()]):
    add_text(slide, skill, .85 + (index % 3) * 1.95, 5.75 + (index // 3) * .5, 1.75, .3, 10, TEXT)
if profile.about_image:
    img = Path(profile.about_image.path)
    if img.exists():
        card(slide, 8.1, 1.1, 4.1, 4.9)
        x, y, w, h = fit_box(img, 8.3, 1.3, 3.7, 4.5)
        slide.shapes.add_picture(str(img), Inches(x), Inches(y), Inches(w), Inches(h))

# Projects: two editable layouts per page
for page in range(0, len(projects), 2):
    slide = prs.slides.add_slide(blank); set_background(slide)
    add_text(slide, "SELECTED WORK", .8, .5, 5, .5, 22, TEXT, True)
    add_text(slide, f"{page + 1:02d} — {min(page + 2, len(projects)):02d}", 10.9, .57, 1.4, .3, 11, PINK, True, align=PP_ALIGN.RIGHT)
    for slot, project in enumerate(projects[page:page + 2]):
        x = .8 + slot * 6.15
        card(slide, x, 1.25, 5.7, 5.45)
        image_path = Path(project.image.path) if project.image else None
        if image_path and image_path.exists():
            ix, iy, iw, ih = fit_box(image_path, x + .25, 1.5, 5.2, 3.35)
            slide.shapes.add_picture(str(image_path), Inches(ix), Inches(iy), Inches(iw), Inches(ih))
        add_text(slide, project.title, x + .28, 5.05, 5.1, .38, 15, TEXT, True)
        add_text(slide, project.category.upper(), x + .28, 5.48, 5.1, .25, 10, PINK, True)
        add_text(slide, clean(project.description), x + .28, 5.83, 5.1, .55, 10, MUTED)

# Contact
slide = prs.slides.add_slide(blank); set_background(slide)
add_text(slide, "LET'S WORK TOGETHER", .8, 1.35, 10.5, .7, 31, TEXT, True)
add_text(slide, "For design enquiries and creative collaborations.", .85, 2.2, 8.5, .4, 15, MUTED)
for n, (label, value) in enumerate((("EMAIL", profile.email), ("PHONE", profile.phone), ("WHATSAPP", profile.whatsapp))):
    y = 3.25 + n * .85
    add_text(slide, label, .85, y, 1.4, .25, 10, PINK, True)
    add_text(slide, value, 2.4, y - .05, 5.5, .35, 15, TEXT)
add_text(slide, "THANK YOU", .85, 6.55, 3, .35, 12, MUTED, True)
prs.save(PPTX_PATH)


W, H = landscape(A4)
pdf = canvas.Canvas(str(PDF_PATH), pagesize=(W, H))
def pdf_bg():
    pdf.setFillColor(HexColor("#" + BG)); pdf.rect(0, 0, W, H, fill=1, stroke=0)
def pdf_text(text, x, y, size, color=TEXT, bold=False):
    pdf.setFont("Helvetica-Bold" if bold else "Helvetica", size); pdf.setFillColor(HexColor("#" + color))
    pdf.drawString(x, y, clean(text))
def pdf_image(path, x, y, w, h):
    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = w / h
    if ratio > box_ratio: iw, ih = w, w / ratio
    else: iw, ih = h * ratio, h
    pdf.drawImage(ImageReader(str(path)), x + (w-iw)/2, y + (h-ih)/2, iw, ih, preserveAspectRatio=True, mask='auto')

pdf_bg(); pdf_text(profile.name.upper(), 55, 370, 29, TEXT, True); pdf_text(profile.title.upper(), 56, 330, 13, PINK, True); pdf_text("PORTFOLIO 2026", 56, 300, 11, MUTED); pdf_text("Bold visuals. Clear ideas. Meaningful design.", 56, 80, 17, TEXT); pdf.showPage()
pdf_bg(); pdf_text("ABOUT ME", 55, 520, 24, TEXT, True); pdf_text(clean(profile.about_text), 55, 465, 14, MUTED); pdf_text(clean(profile.about_detail), 55, 390, 12, MUTED); pdf_text("SKILLS", 55, 280, 11, PINK, True); pdf_text(" • ".join([x.strip() for x in profile.skills.split(',')]), 55, 245, 12, TEXT)
if profile.about_image and Path(profile.about_image.path).exists(): pdf_image(Path(profile.about_image.path), 515, 150, 230, 310)
pdf.showPage()
for project in projects:
    pdf_bg(); pdf_text("SELECTED WORK", 55, 530, 22, TEXT, True); pdf_text(project.category.upper(), 55, 495, 11, PINK, True)
    image_path = Path(project.image.path) if project.image else None
    if image_path and image_path.exists(): pdf_image(image_path, 55, 155, 440, 285)
    pdf_text(project.title, 530, 430, 18, TEXT, True); pdf_text(clean(project.description), 530, 390, 11, MUTED); pdf.showPage()
pdf_bg(); pdf_text("LET'S WORK TOGETHER", 55, 450, 28, TEXT, True); pdf_text(profile.email, 55, 350, 15, TEXT); pdf_text(profile.phone, 55, 305, 15, TEXT); pdf_text(profile.whatsapp, 55, 260, 15, TEXT); pdf.save()
print(PPTX_PATH.resolve()); print(PDF_PATH.resolve())
